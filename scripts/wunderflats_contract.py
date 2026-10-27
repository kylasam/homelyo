#!/usr/bin/env python3
"""
Homelyo Rental Contract — Wunderflats structure (14 clauses)
------------------------------------------------------------
Single source of truth for the bilingual (DE/EN) Homelyo rental contract.

Structure mirrors the signed Wunderflats agreements in Google Drive
(Naeem_Ullah_KE2, Karolina_Faluta_KE4, Sara_Rizzo_KE1) and the Drive sample
`template/template_contract.docx`:

  Mietvertrag (Hauptvertrag)          Anlage 2 – AGB
    § 1 Mietdauer                       § 1 Zusammenleben in der WG
    § 2 Miethöhe                        § 2 Regelungen zur Wohnung
    § 3 Kaution                         § 3 Zahlungsvereinbarungen
    § 4 Zahlungsbedingungen             § 4 Übergabe und Rückgabe
    § 5 Kündigung                       § 5 Schäden an der Wohnung
    § 6 AGB & Anlagen                   § 6 Kündigungsbedingungen
                                        § 7 Schriftform
                                        § 8 Salvatorische Klausel

  + Zahlungsmöglichkeiten, Anlage 1 – Inventar, signatures.

Homelyo rates (override the Wunderflats defaults 245/80/40/30 € keys, 5-week refund):
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
import html as _html
import os

# ----------------------------------------------------------------- Homelyo rates
RATES = {
    "KEY_BUILDING": 680.00,       # Haustürschlüssel
    "KEY_FLAT": 120.00,           # Wohnungsschlüssel
    "KEY_ROOM": 90.00,            # Zimmerschlüssel
    "KEY_MAILBOX": 70.00,         # Briefkastenschlüssel
    "CLEANING_FEE": 70.00,        # einmalige Endreinigung
    "LATE_FEE": 5.00,             # Mahngebühr je Mahnung
    "MINOR_REPAIR_SINGLE": 100.00,  # Kleinreparatur je Einzelfall
    "MINOR_REPAIR_CAP": 350.00,   # Jahreshöchstbetrag
    "MINOR_REPAIR_PCT": 8,        # % der Jahresmiete
    "DAMAGE_TERMINATION": 50.00,  # Schaden > x € unbeglichen -> fristlos
    "DEPOSIT_REFUND_DAYS": 30,
    "NOTICE_MONTHS": 3,           # Kündigungsfrist Mieter (Wunderflats: 1)
    "MIN_TERM_MONTHS": 6,
    "GUEST_DAYS": 5,
}

NUM_DE = {1: "einem", 2: "zwei", 3: "drei", 4: "vier", 5: "fünf", 6: "sechs"}
NUM_EN = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six"}


def eur_de(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def eur_en(v):
    return f"€{v:,.2f}"


# --------------------------------------------------------------- clause content
def build_sections(d):
    """Return the ordered contract sections. `d` is the dict from render_contract_data."""
    R = RATES
    n, m = R["NOTICE_MONTHS"], R["MIN_TERM_MONTHS"]
    rent_w_de = f" (in Worten: {d['RENT_WORDS_DE']} Euro)" if d.get("RENT_WORDS_DE") else ""
    rent_w_en = f" (in words: {d['RENT_WORDS_EN']} euros)" if d.get("RENT_WORDS_EN") else ""
    dep_w_de = f" (in Worten: {d['DEPOSIT_WORDS_DE']} Euro)" if d.get("DEPOSIT_WORDS_DE") else ""
    dep_w_en = f" (in words: {d['DEPOSIT_WORDS_EN']} euros)" if d.get("DEPOSIT_WORDS_EN") else ""
    keys_de = (f"Haustürschlüssel {eur_de(R['KEY_BUILDING'])}, Wohnungsschlüssel {eur_de(R['KEY_FLAT'])}, "
               f"Zimmerschlüssel {eur_de(R['KEY_ROOM'])}, Briefkastenschlüssel {eur_de(R['KEY_MAILBOX'])}")
    keys_en = (f"building entrance key {eur_en(R['KEY_BUILDING'])}, apartment key {eur_en(R['KEY_FLAT'])}, "
               f"room key {eur_en(R['KEY_ROOM'])}, mailbox key {eur_en(R['KEY_MAILBOX'])}")
    bank_de = f"Kontoinhaber: {d['RENT_ACCOUNT_HOLDER']} · IBAN: {d['RENT_IBAN']} · BIC: {d['RENT_BIC']}"
    bank_en = f"Account holder: {d['RENT_ACCOUNT_HOLDER']} · IBAN: {d['RENT_IBAN']} · BIC: {d['RENT_BIC']}"

    main = [
        {
            "de": "§ 1 – Mietdauer", "en": "§ 1 – Term of lease",
            "rows": [
                (f"(1) Das Mietverhältnis beginnt am {d['START_DATE_DE']} und ist bis zum {d['END_DATE_DE']} befristet. "
                 f"Die Mietdauer beträgt mindestens {NUM_DE[m]} ({m}) Monate.",
                 f"(1) The tenancy begins on {d['START_DATE_EN']} and is limited until {d['END_DATE_EN']}. "
                 f"The minimum rental time is {NUM_EN[m]} ({m}) months."),
                (f"(2) Die Vermietung erfolgt zum vorübergehenden Gebrauch gem. § 549 Abs. 2 Nr. 1 BGB: {d['REASON_DE']}.",
                 f"(2) The room is let for temporary use pursuant to § 549 (2) No. 1 BGB: {d['REASON_EN']}."),
                ("(3) Eine stillschweigende Verlängerung des Mietverhältnisses (§ 545 BGB) ist ausgeschlossen.",
                 "(3) Tacit renewal of the tenancy (§ 545 BGB) is excluded."),
            ],
        },
        {
            "de": "§ 2 – Miethöhe", "en": "§ 2 – Rent",
            "rows": [
                (f"(1) Die monatliche Miete für das Zimmer beträgt {d['RENT_AMOUNT_DE']}{rent_w_de}.",
                 f"(1) The monthly rent for the room amounts to {d['RENT_AMOUNT_EN']}{rent_w_en}."),
                ("(2) In der Miete sind die Kosten für Internet (WLAN) und Strom sowie Heiz- und Betriebskosten "
                 "in verbrauchsüblichem Maße enthalten (Pauschalmiete, Fair-Use gem. AGB § 3 Nr. 6).",
                 "(2) The cost for internet (Wi-Fi), electricity, heating and running costs based on customary "
                 "occupancy are included in this price (flat rent, fair use per T&C § 3 no. 6)."),
                ("(3) Der Rundfunkbeitrag (ARD/ZDF) ist nicht enthalten; Anmeldung und Zahlung obliegen den Mietern (AGB § 2 Nr. 3 d).",
                 "(3) The public broadcasting fee (ARD/ZDF) is not included; registration and payment are the tenants' duty (T&C § 2 no. 3 d)."),
                (f"(4) Bei Auszug wird eine einmalige Endreinigungspauschale von {eur_de(R['CLEANING_FEE'])} fällig.",
                 f"(4) A one-time final cleaning fee of {eur_en(R['CLEANING_FEE'])} is due upon move-out."),
            ],
        },
        {
            "de": "§ 3 – Kaution", "en": "§ 3 – Security deposit",
            "rows": [
                (f"Der Mieter leistet zur Sicherung aller Ansprüche des Vermieters eine Kaution in Höhe von "
                 f"{d['DEPOSIT_AMOUNT_DE']}{dep_w_de}. Die Kaution wird getrennt vom Vermögen des Vermieters angelegt (§ 551 Abs. 3 BGB).",
                 f"A security deposit of {d['DEPOSIT_AMOUNT_EN']}{dep_w_en} shall secure compliance with the terms and "
                 f"conditions of this agreement. It is held separately from the landlord's assets (§ 551 (3) BGB)."),
            ],
        },
        {
            "de": "§ 4 – Zahlungsbedingungen", "en": "§ 4 – Payment",
            "rows": [
                (f"(1) Die Kaution ist bis zum {d['DEPOSIT_DUE_DE']} per Banküberweisung zu leisten, die erste Miete ist "
                 f"vor Übergabe des Mietobjektes zu entrichten.",
                 f"(1) The security deposit is to be paid by bank transfer before {d['DEPOSIT_DUE_EN']}. "
                 f"The first month's rent must be paid before handover of the keys."),
                ("(2) Die Gesamtmiete ist im Voraus und kostenfrei bis zum 3. Werktag eines jeden Monats zu überweisen.",
                 "(2) Rent must be paid in advance and free of charges by the third working day of each month."),
                (f"(3) {bank_de}.", f"(3) {bank_en}."),
            ],
        },
        {
            "de": "§ 5 – Kündigung", "en": "§ 5 – Termination",
            "rows": [
                (f"(1) Dem Mieter wird ein verkürztes Kündigungsrecht eingeräumt: Er kann das Mietverhältnis mit einer Frist "
                 f"von {NUM_DE[n]} ({n}) Monaten zum Ende eines Kalendermonats schriftlich kündigen, frühestens jedoch "
                 f"zum Ablauf der Mindestmietdauer von {m} Monaten.",
                 f"(1) The tenant is granted a simplified right of termination: the tenant may terminate the tenancy in "
                 f"writing with {NUM_EN[n]} ({n}) months' notice to the end of a calendar month, at the earliest "
                 f"effective at the end of the {m}-month minimum term."),
                ("(2) Das Recht zur außerordentlichen fristlosen Kündigung aus wichtigem Grund (§§ 543, 569 BGB, AGB § 6) bleibt unberührt.",
                 "(2) The right to extraordinary termination for good cause (§§ 543, 569 BGB, T&C § 6) remains unaffected."),
            ],
        },
        {
            "de": "§ 6 – Allgemeine Geschäftsbedingungen & Anlagen", "en": "§ 6 – Terms and conditions & attachments",
            "rows": [
                ("Die angehängten Anlagen – Anlage 1 (Inventar der Mietsache) und Anlage 2 (Allgemeine Geschäftsbedingungen) – "
                 "werden mit der Unterschrift als Vertragsbestandteil akzeptiert und sind von Mieter und Vermieter einzuhalten. "
                 "Übergabeprotokoll, Schlüsselübergabe und Wohnungsgeberbestätigung (§ 19 BMG) werden bei Einzug ausgehändigt.",
                 "The attachments – Attachment 1 (Inventory) and Attachment 2 (General Terms and Conditions) – come into "
                 "effect upon signature of this agreement and are to be upheld by both landlord and tenant. Handover "
                 "protocol, key acknowledgement and landlord confirmation (§ 19 BMG) are provided at move-in."),
            ],
        },
    ]

    agb = [
        {
            "de": "§ 1 Zusammenleben in der WG / mit den Nachbarn",
            "en": "§ 1 Cohabitation in a flat share / with the neighbours",
            "rows": [
                ("1. Rücksicht: Die WG hat das Zusammenleben untereinander mit Rücksicht auf die Interessen und zur "
                 "Zufriedenheit aller Beteiligten fair und eigenständig zu regeln.",
                 "1. Regard: All flat members organise their cohabitation fairly and with consideration of the interests of all members."),
                ("2. Lärm: a) Vermeidbarer Lärm in Wohnung, Haus, Hof und Grundstück hat zu unterbleiben. b) Ruhezeiten "
                 "13:00–15:00 Uhr und 22:00–08:00 Uhr; an Sonn- und Feiertagen ist Lärm komplett zu unterlassen; Soundanlagen/TV "
                 "auf Zimmerlautstärke. c) Musizieren ist 13:00–15:00 Uhr und 19:00–08:00 Uhr untersagt, sonst max. zwei Stunden täglich.",
                 "2. Noise: a) Avoidable noise in the flat, building, courtyard and property is prohibited. b) Quiet hours "
                 "1:00–3:00 pm and 10:00 pm–8:00 am; no noise on Sundays and public holidays; stereos/TV at room volume. "
                 "c) Instruments are forbidden 1:00–3:00 pm and 7:00 pm–8:00 am, otherwise max. two hours per day."),
                ("3. Partys: Störende Partys sind untersagt. Beim Zusammensitzen ist die Lautstärke spätestens um 22:00 Uhr "
                 "auf Zimmerlautstärke zu beschränken.",
                 "3. Parties: Noisy parties are prohibited. Gatherings must be reduced to room volume by 10:00 pm at the latest."),
                ("4. Reinigung: a) Die Mieter reinigen Zimmer und Gemeinschaftsräume (Bad, WC, Küche, Flur, Balkon) eigenverantwortlich. "
                 "b) Erfolgt trotz Fristsetzung keine Reinigung, kann der Vermieter eine professionelle Reinigung auf Kosten aller "
                 "Mieter zu gleichen Teilen beauftragen. c) Bei mehrmaliger Nichtbeteiligung ist nach Abmahnung eine fristlose Kündigung möglich.",
                 "4. Cleaning: a) Tenants clean their room and shared areas (bath, WC, kitchen, hall, balcony) themselves. "
                 "b) If cleaning is not done despite a deadline, the landlord may order professional cleaning at the equal expense "
                 "of all tenants. c) Repeated non-participation entitles the landlord to terminate without notice after a warning."),
                ("5. Rauchverbot / Drogen: a) Rauchen (inkl. E-Zigaretten) ist in der gesamten Wohnung, auf dem Balkon und im Haus untersagt. "
                 "b) Besitz und Konsum illegaler Drogen auf dem gesamten Grundstück sind verboten, werden polizeilich angezeigt und "
                 "berechtigen zur außerordentlichen Kündigung.",
                 "5. Smoking / drugs: a) Smoking (incl. e-cigarettes) is forbidden in the entire flat, balcony and building. "
                 "b) Possession or use of illegal drugs anywhere on the property is prohibited, will be reported to the police "
                 "and entitles the landlord to immediate termination."),
                ("6. Haustiere: Die Haltung von Haustieren ist grundsätzlich verboten.",
                 "6. Pets: Keeping pets is generally forbidden."),
                (f"7. Beherbergung anderer Personen: a) Eine Überlassung der Mieträume an Dritte (entgeltlich oder unentgeltlich) "
                 f"sowie die dauerhafte Aufnahme eines Lebensgefährten sind nicht gestattet. b) Besucher dürfen an höchstens "
                 f"{R['GUEST_DAYS']} Tagen pro Monat und nur bei Anwesenheit des Mieters beherbergt werden; andernfalls ist nach "
                 f"Abmahnung eine fristlose Kündigung zulässig.",
                 f"7. Lodging of other persons: a) Letting the room to third parties (paid or free) or permanently "
                 f"accommodating a partner is prohibited. b) Visitors may stay at most {R['GUEST_DAYS']} days per month and only "
                 f"while the tenant is present; otherwise the landlord may terminate without notice after a warning."),
            ],
        },
        {
            "de": "§ 2 Regelungen zur Wohnung", "en": "§ 2 Regulations concerning the flat",
            "rows": [
                ("1. Nutzung: Die Vermietung erfolgt nur zu Wohnzwecken; gewerbliche Nutzung ist nicht gestattet.",
                 "1. Use: The room is let for residential purposes only; commercial use is prohibited."),
                ("2. Lüften und Heizen: Mindestens zweimal täglich kurz und vollständig lüften (Stoßlüften), auch im Winter. "
                 "Heizkörper beim Lüften und bei längerer Abwesenheit zudrehen.",
                 "2. Ventilation and heating: Air out at least twice daily by fully opening windows briefly, also in winter. "
                 "Turn radiators off when airing or when absent for longer."),
                ("3. Internet und Telefon: a) Eine Telefonanlage darf nicht installiert werden. b) Der Vermieter stellt WLAN zur Verfügung. "
                 "c) Der Mieter trägt die volle zivil- und strafrechtliche Verantwortung für seine Nutzung; Filesharing (P2P/Torrent) "
                 "und illegale Downloads sind verboten. Der Mieter stellt den Vermieter von Abmahnungen, Schadensersatz- und Anwaltskosten "
                 "frei; bei Verstößen kann der Zugang gesperrt und die Mieterdaten nach gesetzlichen Vorgaben an Behörden weitergegeben werden. "
                 "d) Die Anmeldung des Rundfunkbeitrags ist vom Mieter vorzunehmen; eine Lösung ist innerhalb der WG zu finden.",
                 "3. Internet and phone: a) No private telephone system may be installed. b) The landlord provides Wi-Fi. "
                 "c) The tenant bears full civil and criminal responsibility for their use; file sharing (P2P/torrent) and illegal "
                 "downloads are prohibited. The tenant indemnifies the landlord against warning notices, damages and legal fees; "
                 "in case of violations access may be blocked and tenant data passed to authorities as required by law. "
                 "d) Registration for the broadcasting fee is the tenant's duty and is to be settled within the flat share."),
                ("4. Bilder und Plakate: Nägel, Schrauben oder Dübel sind untersagt. Schilder/Aufschriften in Gemeinschaftsräumen "
                 "sind verboten. Der Vermieter beschriftet das Klingelschild einheitlich; eigene Schilder nur mit Genehmigung.",
                 "4. Pictures and posters: Nails, screws or wall plugs are forbidden. Signs in shared rooms are prohibited. "
                 "The landlord labels the doorbell uniformly; personal signs only with permission."),
                ("5. Obhut und Sorgfalt: Keine Abfälle, Essensreste, Fette oder Hygieneartikel in Abflüsse. Böden sind zu pflegen "
                 "und vor Kratzern zu schützen.",
                 "5. Duty of care: No waste, food, grease or sanitary items down the drains. Floors must be cared for and protected from scratches."),
                ("6. Müllentsorgung: Müll nur in die vorgesehenen Tonnen, konsequent getrennt (BSR). Sperr- und Sondermüll sind gesondert zu entsorgen.",
                 "6. Waste: Dispose of rubbish only in the designated, separated bins (BSR). Bulky and hazardous waste separately."),
                ("7. Bauliche Veränderungen: Untersagt. Satellitenanlagen nur mit schriftlicher Genehmigung. Notwendige Erhaltungsmaßnahmen "
                 "sind nach Ankündigung zu dulden.",
                 "7. Modifications: Prohibited. Satellite dishes only with written permission. Necessary maintenance works must be tolerated after notice."),
                ("8. Sicherheit: Keine Lagerung giftiger oder brennbarer Stoffe. Haus-, Keller- und Hoftüren 22:00–06:00 Uhr geschlossen halten; "
                 "Fluchtwege freihalten. Grillen auf dem Balkon und offenes Feuer sind untersagt. Bei Gasgeruch: kein Licht/Schalter, "
                 "Fenster öffnen, Haupthahn schließen, Versorger und Vermieter informieren.",
                 "8. Safety: No storage of toxic or flammable substances. Keep building, basement and yard doors closed 10:00 pm–6:00 am; "
                 "keep escape routes clear. Balcony grilling and open fire are forbidden. If gas is smelled: no light/switches, open "
                 "windows, close main valve, inform supplier and landlord."),
                ("9. Betreten durch den Vermieter: a) Werktags 09:00–12:00 und 15:00–20:00 Uhr nach Ankündigung, auch mit Fachleuten "
                 "oder Mietinteressenten; in den letzten vier Wochen mind. fünf Besichtigungen. b) Bei konkreter Gefahr jederzeit. "
                 "c) Der Mieter stellt dies auch bei Abwesenheit sicher. d) Zur Kontrolle der Reinigung der Gemeinschaftsflächen.",
                 "9. Landlord access: a) Working days 9:00 am–12:00 pm and 3:00–8:00 pm after notice, also with technicians or "
                 "prospective tenants; at least five viewings during the last four weeks. b) Any time in case of imminent danger. "
                 "c) The tenant ensures this also during absence. d) To check cleaning of shared areas."),
            ],
        },
        {
            "de": "§ 3 Zahlungsvereinbarungen", "en": "§ 3 Payment agreement",
            "rows": [
                ("1. Rechtzeitigkeit: Maßgeblich ist das Eingangsdatum, nicht das Absendedatum.",
                 "1. Punctuality: The date of receipt, not the date of submission, is decisive."),
                ("2. Kaution: Sichert alle Ansprüche des Vermieters (Vertragserfüllung, Schadensersatz, nicht zurückgegebene Schlüssel). "
                 "Zahlung per Überweisung, Verwendungszweck: „Kaution, Vor- und Nachname“. Eine Aufrechnung mit dem Rückzahlungsanspruch "
                 "vor Mietende ist ausgeschlossen.",
                 "2. Deposit: Secures all landlord claims (performance, damages, unreturned keys). Paid by bank transfer, reference "
                 "\"Deposit, first and last name\". Offsetting the refund claim before the end of the tenancy is excluded."),
                ("3. Mietzahlung: Im Voraus und kostenfrei bis zum 3. Werktag eines Monats, Verwendungszweck: „Vor- und Nachname, Monat“.",
                 "3. Rent: In advance and free of charges by the 3rd working day of each month, reference \"first and last name, month\"."),
                ("4. Zahlungsarten: Banküberweisung; Barzahlung nur gegen Quittung. Gebühren der Zahlungsart trägt der Mieter.",
                 "4. Payment methods: Bank transfer; cash only against receipt. Any fees of the payment method are borne by the tenant."),
                (f"5. Rückerstattung der Kaution: Abrechnung innerhalb von {R['DEPOSIT_REFUND_DAYS']} Tagen nach Mietende und Rückgabe, "
                 f"abzüglich begründeter Gegenansprüche, per Überweisung. Der Mieter teilt seine Kontodaten spätestens beim Auszug mit.",
                 f"5. Deposit refund: Settled within {R['DEPOSIT_REFUND_DAYS']} days after the end of the tenancy and return of the room, "
                 f"less justified claims, by bank transfer. The tenant provides bank details at move-out at the latest."),
                ("6. Fair Usage: Pauschalmiete ohne jährliche Nebenkostenabrechnung. Übersteigen die Nebenkosten (§ 2 BetrKV) den Durchschnitt "
                 "um mehr als 25 %, kann der Mehrbetrag nach billigem Ermessen umgelegt werden (Gegenbeweis bleibt möglich; Mieter haften "
                 "als Gesamtschuldner). Geräte mit hohem Energiebedarf sind anzuzeigen.",
                 "6. Fair usage: Flat rent without annual utility statement. If utilities (§ 2 BetrKV) exceed the average by more than 25 %, "
                 "the excess may be charged at reasonable discretion (counter-proof remains possible; tenants are jointly liable). "
                 "High-consumption appliances must be reported."),
                (f"7. Mahngebühren: {eur_de(R['LATE_FEE'])} je Mahnung.",
                 f"7. Late fees: {eur_en(R['LATE_FEE'])} per reminder notice."),
                ("8. Mieterhöhung: Nach Maßgabe der gesetzlichen Bestimmungen.",
                 "8. Rent increase: Within the limits of statutory law."),
                ("9. Aufrechnung: Nur nach Anzeige mindestens einen Monat vor Fälligkeit; ausgenommen unbestrittene oder rechtskräftige "
                 "Forderungen. §§ 536, 536a BGB bleiben unberührt.",
                 "9. Set-off: Only after notice at least one month before the due date; except undisputed or final claims. "
                 "§§ 536, 536a BGB remain unaffected."),
            ],
        },
        {
            "de": "§ 4 Übergabe und Rückgabe",
            "en": "§ 4 Handover and return of the property",
            "rows": [
                ("1. Zeitpunkt: Nach Eingang von Kaution und erster Miete, frühestens zum Mietbeginn.",
                 "1. Timing: After receipt of deposit and first rent, at the earliest on the start date."),
                ("2. Zustand bei Einzug: Es wird ein Übergabeprotokoll erstellt; dokumentierte Schäden/Abnutzung sind bei Auszug im selben Zustand zu hinterlassen.",
                 "2. Condition at move-in: A handover protocol is signed; documented damage/wear must be left in the same condition at move-out."),
                (f"3. Schlüssel: a) Anzahl und Art werden in der Schlüsselübergabe festgehalten. b) Bei Verlust oder Bruch trägt der Mieter: "
                 f"{keys_de}. Der Nachweis eines geringeren Schadens bleibt unbenommen. c) Zusätzlich haftet der Mieter für notwendige "
                 f"Schlossänderungen (Haus-, Wohnungs-, Zimmertür, Briefkasten / Schließanlage).",
                 f"3. Keys: a) Number and type are recorded in the key acknowledgement. b) For loss or breakage the tenant pays: "
                 f"{keys_en}. Proof of a lower damage remains possible. c) The tenant is additionally liable for necessary lock "
                 f"changes (building, apartment, room door, mailbox / master system)."),
                ("4. Rückgabe: a) Am letzten Tag der Mietzeit bis 11:00 Uhr; Termin wird vereinbart. b) Vollständig geräumt und besenrein "
                 "mit allen Schlüsseln (Endreinigung über die Pauschale gem. § 2 Abs. 4 Mietvertrag). c) Veränderungen sind auf Verlangen "
                 "zurückzubauen. d) Zurückgelassene Gegenstände werden auf Kosten des Mieters entsorgt.",
                 "4. Return: a) On the last day of the tenancy by 11:00 am; an appointment is arranged. b) Fully cleared and broom-clean "
                 "with all keys (final cleaning via the fee per § 2 (4) of the agreement). c) Modifications must be reversed on request. "
                 "d) Items left behind are disposed of at the tenant's expense."),
            ],
        },
        {
            "de": "§ 5 Schäden an der Wohnung", "en": "§ 5 Damages to the property",
            "rows": [
                ("1. Anzeigepflicht: Mängel unverzüglich melden (auch Inventar gem. Anlage 1); der Vermieter erhält 4 Wochen zur Abhilfe. "
                 "Aufwendungsersatz nur nach Gelegenheit zur Beseitigung, außer bei Gefahr im Verzug.",
                 "1. Reporting: Report defects immediately (incl. inventory per Attachment 1); the landlord has 4 weeks to remedy. "
                 "Reimbursement only after the landlord had the chance to repair, except in emergencies."),
                ("2. Haftung: a) Der Mieter haftet für Schäden durch Verletzung von Sorgfalts- und Anzeigepflichten (u. a. unzureichendes "
                 "Lüften/Heizen) sowie b) für Personen seines Hausstands und Besucher. c) Beweislast im Zimmer und in Gemeinschaftsräumen "
                 "beim Mieter. d) Verschuldensunabhängige Haftung des Vermieters für anfängliche Mängel ist ausgeschlossen. e) Nach "
                 "Vertragsschluss haftet der Vermieter nur für Vorsatz und grobe Fahrlässigkeit (Ausnahmen: Kardinalpflichten, Leben, Körper, Gesundheit).",
                 "2. Liability: a) The tenant is liable for damage from breach of duties of care/reporting (incl. insufficient airing/heating) "
                 "and b) for household members and visitors. c) Burden of proof for the room and shared areas lies with the tenant. "
                 "d) Strict liability of the landlord for initial defects is excluded. e) After signing, the landlord is liable only for "
                 "intent and gross negligence (exceptions: essential obligations, life, body, health)."),
                ("3. Schönheitsreparaturen: Nicht geschuldet, außer bei übermäßiger Abnutzung durch den Mieter.",
                 "3. Cosmetic repairs: Not owed, except in case of excessive wear caused by the tenant."),
                ("4. Schadensbeseitigung: Fachgerecht in Abstimmung mit dem Vermieter; verlorene/beschädigte Gegenstände sind gleichwertig zu ersetzen.",
                 "4. Repair of damage: Professionally and in coordination with the landlord; lost/damaged items must be replaced like-for-like."),
                (f"5. Kleinreparaturen: Kosten für Teile im ständigen Gebrauch (Leuchtmittel, Fenster-/Türgriffe, Armaturen, Schalter) trägt "
                 f"der Mieter bis {eur_de(R['MINOR_REPAIR_SINGLE'])} je Einzelfall, jährlich max. {R['MINOR_REPAIR_PCT']} % der Jahresmiete, "
                 f"höchstens {eur_de(R['MINOR_REPAIR_CAP'])}. Die Auftragsvergabe erfolgt durch den Vermieter.",
                 f"5. Minor repairs: Costs for parts in constant use (bulbs, window/door handles, fittings, switches) are borne by the "
                 f"tenant up to {eur_en(R['MINOR_REPAIR_SINGLE'])} per case, annually max. {R['MINOR_REPAIR_PCT']} % of the annual rent, "
                 f"at most {eur_en(R['MINOR_REPAIR_CAP'])}. Contractors are commissioned by the landlord."),
            ],
        },
        {
            "de": "§ 6 Kündigungsbedingungen", "en": "§ 6 Conditions of termination",
            "rows": [
                (f"1. Kündigungsfrist: Gem. § 5 Mietvertrag ({n} Monate zum Monatsende, frühestens zum Ende der Mindestmietdauer). "
                 f"Maßgeblich ist der Zugang des Kündigungsschreibens.",
                 f"1. Notice period: Per § 5 of the agreement ({n} months to month-end, at the earliest at the end of the minimum term). "
                 f"The date of receipt of the notice is decisive."),
                ("2. Nachmieter: Der Mieter kann einen Nachmieter vorschlagen; der Vermieter prüft umgehend und kann ohne Angabe von Gründen "
                 "ablehnen. Bei Annahme endet die Zahlungspflicht ab Übergabe an den Nachmieter; andernfalls bleibt der Mieter bis Fristablauf zahlungspflichtig.",
                 "2. Replacement tenant: The tenant may propose a successor; the landlord reviews promptly and may reject without reason. "
                 "If accepted, payment obligations end upon handover to the successor; otherwise the tenant pays until the notice expires."),
                ("3. Stillschweigende Verlängerung: Ausgeschlossen; § 545 BGB findet keine Anwendung.",
                 "3. Tacit renewal: Excluded; § 545 BGB does not apply."),
                (f"4. Fristlose Kündigung: a) Der Mieter haftet für Mietausfall bis zum Ende der ordentlichen Frist. b) Gründe u. a.: Verstoß "
                 f"gegen die AGB trotz Abmahnung; vom Mieter verursachte Schäden über {eur_de(R['DAMAGE_TERMINATION'])}, die trotz Fristsetzung "
                 f"nicht ausgeglichen werden; unerlaubte Überlassung an Dritte; Unterlassen der Reinigung trotz Abmahnung; 2 Abmahnungen für "
                 f"denselben Tatbestand oder 3 insgesamt; Zahlungsverzug über zwei Termine (§ 543 Abs. 2 BGB). Im Übrigen gelten die gesetzlichen Bestimmungen.",
                 f"4. Termination without notice: a) The tenant is liable for loss of rent until the ordinary notice period ends. b) Grounds "
                 f"include: breach of T&C despite warning; tenant-caused damage over {eur_en(R['DAMAGE_TERMINATION'])} unpaid after a deadline; "
                 f"unauthorised third-party lodging; failure to clean despite warning; 2 warnings for the same matter or 3 in total; rent "
                 f"arrears over two due dates (§ 543 (2) BGB). Otherwise statutory law applies."),
            ],
        },
        {
            "de": "§ 7 Schriftform", "en": "§ 7 Written form",
            "rows": [
                ("Mündliche Nebenabreden bestehen nicht. Änderungen und Ergänzungen bedürfen der Schriftform; dies gilt auch für die "
                 "Aufhebung dieses Schriftformerfordernisses.",
                 "There are no verbal side agreements. Changes and additions must be in writing; this also applies to waiving this written-form requirement."),
            ],
        },
        {
            "de": "§ 8 Salvatorische Klausel & Gerichtsstand", "en": "§ 8 Severability & jurisdiction",
            "rows": [
                ("Sollten einzelne Bestimmungen unwirksam oder undurchführbar sein oder werden, bleibt die Wirksamkeit der übrigen "
                 "unberührt. An ihre Stelle tritt die Regelung, die dem wirtschaftlichen Zweck am nächsten kommt; Grundlage ist das BGB. "
                 "Erfüllungsort und Gerichtsstand ist Berlin. Die deutsche Fassung ist allein rechtsverbindlich.",
                 "Should individual provisions be or become invalid or unenforceable, the remaining provisions stay valid. They are replaced "
                 "by the provision closest to the economic purpose, based on the German Civil Code (BGB). Place of performance and "
                 "jurisdiction is Berlin. Only the German version is legally binding."),
            ],
        },
    ]

    inventory = [
        ("Ausstattung im Zimmer: Bett inkl. Lattenrost, Matratze & Schoner; Kleiderschrank/Kommode; Schreibtisch & Bürostuhl; "
         "Spiegel, Beleuchtung, Vorhänge; Bettwäsche (Kissen, Decke, Laken, Bezüge). "
         f"Zimmerspezifisch: {d['ROOM_FEATURES_DE']}.",
         "In the room: bed with slatted frame, mattress & protector; wardrobe/dresser; desk & office chair; mirror, lighting, "
         f"curtains; bed linen (pillow, duvet, sheet, covers). Room-specific: {d['ROOM_FEATURES_EN']}."),
        ("Gemeinschaftsausstattung: Waschmaschine, Kühlschrank, Herd/Induktionskochfeld, Backofen, Mikrowelle, Kaffeemaschine, "
         "Wasserkocher, Staubsauger, Geschirr & Kochutensilien.",
         "Shared: washing machine, fridge, stove/induction hob, oven, microwave, coffee maker, kettle, vacuum cleaner, tableware & cookware."),
        (f"Wichtig: Schäden sind gem. AGB § 5 Nr. 1 unverzüglich zu melden; vom Mieter verursachte Schäden trägt er gem. AGB § 5 Nr. 2 a) "
         f"in vollem Maße; Kleinreparaturen bis {eur_de(R['MINOR_REPAIR_SINGLE'])} (inkl. Leuchtmittel) gem. AGB § 5 Nr. 5. Ersatz für "
         f"nicht mitvermietete Gegenstände kann nicht verlangt werden.",
         f"Important: Damage must be reported immediately (T&C § 5 no. 1); tenant-caused damage is fully borne by the tenant (T&C § 5 "
         f"no. 2 a); minor repairs up to {eur_en(R['MINOR_REPAIR_SINGLE'])} (incl. bulbs) per T&C § 5 no. 5. No claims for items not included."),
    ]
    return main, agb, inventory


AGB_PREAMBLE = (
    "Der Mieter erkennt die Allgemeinen Geschäftsbedingungen (AGB) als für ihn verbindlich an. Ein Verstoß gegen die AGB ist ein "
    "vertragswidriger Gebrauch des Mietgegenstandes; in schwerwiegenden Fällen kann der Vermieter ohne Einhaltung der Kündigungsfrist "
    "kündigen. Für Schäden aus der Verletzung der AGB oder der Meldepflichten ist der Mieter ersatzpflichtig. Die AGB gelten für alle "
    "Bewohner der Wohngemeinschaft.",
    "The tenant accepts the general terms and conditions as binding. A violation is a breach of contract; in severe cases the landlord "
    "may terminate without notice. The tenant must compensate for damage caused by violating these terms or the reporting duties. "
    "The terms apply to all residents of the flat share equally.",
)


def _parties(d):
    de = (f"Vermieter: {d['LANDLORD_SIGN_NAME']}, {d['LANDLORD_ADDRESS_DE']}, {d['LANDLORD_CONTACT']} (nachfolgend „Vermieter“)\n"
          f"Mieter: {d['TENANT_NAME']}, {d['TENANT_ADDRESS_DE']}, {d['TENANT_EMAIL']} | {d['TENANT_PHONE']} (nachfolgend „Mieter“)\n"
          f"über ein Zimmer (Zimmer-ID {d['ROOM_CODE']}, {d['ROOM_TITLE_DE']}, ca. {d['ROOM_SIZE']} m²) in der Wohnung: "
          f"Lichtenrader Damm 80, 12305 Berlin, Etage {d['FLOOR']}, Wohnungsnummer WE-11, Zimmernummer {d['ROOM_COUNT']}, Deutschland. "
          f"Flur, Küche, Bad und Balkon/Terrasse dürfen mitbenutzt werden.")
    en = (f"Landlord: {d['LANDLORD_SIGN_NAME']}, {d['LANDLORD_ADDRESS_EN']}, {d['LANDLORD_CONTACT']} (\"landlord\")\n"
          f"Tenant: {d['TENANT_NAME']}, {d['TENANT_ADDRESS_EN']}, {d['TENANT_EMAIL']} | {d['TENANT_PHONE']} (\"tenant\")\n"
          f"concerning a room (Room ID {d['ROOM_CODE']}, {d['ROOM_TITLE_EN']}, approx. {d['ROOM_SIZE']} m²) in the apartment: "
          f"Lichtenrader Damm 80, 12305 Berlin, floor {d['FLOOR']}, flat no. WE-11, room no. {d['ROOM_COUNT']}, Germany. "
          f"Hallway, kitchen, bathroom and balcony/terrace may be shared.")
    return de, en


# ------------------------------------------------------------------ DOCX output
GREY = RGBColor(71, 85, 105)


def _heading(doc, de, en, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(de + "\n"); r.bold = True; r.font.size = Pt(size)
    r = p.add_run(en); r.font.size = Pt(size - 3); r.font.color.rgb = GREY


def _table(doc, rows, head=None):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    if head:
        cells = t.add_row().cells
        for c, txt in zip(cells, head):
            r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.size = Pt(9.5)
    for de, en in rows:
        cells = t.add_row().cells
        cells[0].paragraphs[0].add_run(de).font.size = Pt(8.5)
        r = cells[1].paragraphs[0].add_run(en); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    for row in t.rows:
        row.cells[0].width = Inches(3.55); row.cells[1].width = Inches(3.55)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _signatures(doc, d):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].paragraphs[0].add_run(
        f"Ort, Datum / Place, date: Berlin, {d['LANDLORD_SIGN_DATE']}\n\n\n_____________________________\n"
        f"{d['LANDLORD_SIGN_NAME']}\nUnterschrift Vermieter / Landlord signature")
    t.rows[0].cells[1].paragraphs[0].add_run(
        f"Ort, Datum / Place, date: {d['TENANT_SIGN_DATE']}\n\n\n_____________________________\n"
        f"{d['TENANT_SIGN_NAME']}\nUnterschrift Mieter / Tenant signature")


def render_docx(d, output_path, logo_path=None):
    main, agb, inventory = build_sections(d)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(0.55)

    # Cover + welcome
    if logo_path and os.path.exists(logo_path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.8))
    _heading(doc, "Mietvertrag", "Rental Agreement", 24)
    _table(doc, [
        ("Mieter / Tenant", d["TENANT_NAME"]),
        ("Mietobjekt / Room", f"Zimmer {d['ROOM_CODE']} ({d['ROOM_TITLE_DE']} / {d['ROOM_TITLE_EN']})"),
        ("Liegenschaft / Address", "Lichtenrader Damm 80, 12305 Berlin"),
        ("Mietzeit / Rental term", f"{d['START_DATE_DE']} – {d['END_DATE_DE']}"),
        ("Datum / Date", d["CONTRACT_DATE"]),
    ])
    _table(doc, [(
        f"Lieber Mieter, {d['TENANT_NAME']},\nim Folgenden erhältst Du alle wichtigen Informationen rund um Deinen Mietvertrag:\n"
        "• Mietvertrag\n• Zahlungsmöglichkeiten\n• Anlage 1 – Inventar der Mietsache\n• Anlage 2 – Allgemeine Geschäftsbedingungen",
        f"Dear tenant, {d['TENANT_NAME']},\nbelow you will find all important information regarding your rental agreement:\n"
        "• Rental agreement\n• Payment options\n• Attachment 1 – Inventory\n• Attachment 2 – General terms and conditions",
    )])
    doc.add_page_break()

    # Main agreement § 1 – § 6
    _heading(doc, "Mietvertrag zwischen", "Rental agreement between", 13)
    _table(doc, [_parties(d)])
    for sec in main:
        _table(doc, sec["rows"], head=(sec["de"], sec["en"]))
    _signatures(doc, d)
    p = doc.add_paragraph("Hinweis: Die deutsche Fassung ist allein rechtsverbindlich. / Please note: the English translation is for information only; the German version is legally binding.")
    p.runs[0].font.size = Pt(7.5); p.runs[0].italic = True
    doc.add_page_break()

    # Payment options
    _heading(doc, "Zahlungsmöglichkeiten", "Payment options", 13)
    _table(doc, [
        ("Kontoinhaber / Account holder", d["RENT_ACCOUNT_HOLDER"]),
        ("IBAN", d["RENT_IBAN"]), ("BIC / SWIFT", d["RENT_BIC"]), ("Bank", "N26 Bank AG, Berlin"),
        ("Verwendungszweck Miete / Rent reference", f"Miete [Monat Jahr] {d['TENANT_NAME']}"),
        ("Verwendungszweck Kaution / Deposit reference", f"Kaution {d['TENANT_NAME']} Zimmer {d['ROOM_CODE']}"),
    ])
    # Anlage 1
    _heading(doc, "Anlage 1 – Inventar der Mietsache", "Attachment 1 – Inventory", 13)
    _table(doc, inventory)
    doc.add_page_break()

    # Anlage 2 AGB § 1 – § 8
    _heading(doc, "Anlage 2 – Allgemeine Geschäftsbedingungen", "Attachment 2 – General Terms and Conditions", 13)
    _table(doc, [AGB_PREAMBLE])
    for sec in agb:
        _table(doc, sec["rows"], head=(sec["de"], sec["en"]))
    _signatures(doc, d)

    doc.save(output_path)
    return output_path


# ------------------------------------------------------------------ HTML output
def render_html(d, logo_tag=""):
    main, agb, inventory = build_sections(d)
    e = lambda s: _html.escape(str(s)).replace("\n", "<br>")

    def table(rows, head=None):
        h = f'<tr><th>{e(head[0])}</th><th>{e(head[1])}</th></tr>' if head else ""
        return '<table>' + h + "".join(f"<tr><td>{e(a)}</td><td class='en'>{e(b)}</td></tr>" for a, b in rows) + "</table>"

    sig = (f"<table class='sig'><tr><td>Berlin, {e(d['LANDLORD_SIGN_DATE'])}<br><br><br>______________________<br>"
           f"{e(d['LANDLORD_SIGN_NAME'])}<br>Vermieter / Landlord</td><td>{e(d['TENANT_SIGN_DATE'])}<br><br><br>"
           f"______________________<br>{e(d['TENANT_SIGN_NAME'])}<br>Mieter / Tenant</td></tr></table>")
    body = [
        f"<div class='page'>{logo_tag}<h1>Mietvertrag<br><small>Rental Agreement</small></h1>",
        table([("Mieter / Tenant", d["TENANT_NAME"]), ("Mietobjekt / Room", f"Zimmer {d['ROOM_CODE']} ({d['ROOM_TITLE_DE']})"),
               ("Liegenschaft / Address", "Lichtenrader Damm 80, 12305 Berlin"),
               ("Mietzeit / Rental term", f"{d['START_DATE_DE']} – {d['END_DATE_DE']}"), ("Datum / Date", d["CONTRACT_DATE"])]),
        "</div><div class='page'><h2>Mietvertrag zwischen / Rental agreement between</h2>", table([_parties(d)]),
        "".join(table(s["rows"], (s["de"], s["en"])) for s in main), sig,
        "<p class='note'>Die deutsche Fassung ist allein rechtsverbindlich. / The German version is legally binding.</p></div>",
        "<div class='page'><h2>Zahlungsmöglichkeiten / Payment options</h2>",
        table([("Kontoinhaber / Account holder", d["RENT_ACCOUNT_HOLDER"]), ("IBAN", d["RENT_IBAN"]), ("BIC", d["RENT_BIC"]),
               ("Bank", "N26 Bank AG, Berlin")]),
        "<h2>Anlage 1 – Inventar / Attachment 1 – Inventory</h2>", table(inventory), "</div>",
        "<div class='page'><h2>Anlage 2 – Allgemeine Geschäftsbedingungen / General Terms and Conditions</h2>",
        table([AGB_PREAMBLE]), "".join(table(s["rows"], (s["de"], s["en"])) for s in agb), sig, "</div>",
    ]
    css = ("@page{size:A4;margin:11mm 13mm}body{font-family:-apple-system,Helvetica,Arial,sans-serif;font-size:8.4pt;"
           "line-height:1.34;color:#111827}h1{text-align:center;font-size:22pt}h1 small{font-size:12pt;color:#475569}"
           "h2{font-size:11pt;border-bottom:1px solid #e2e8f0;padding-bottom:4px}table{width:100%;border-collapse:collapse;"
           "margin-bottom:8px}td,th{width:50%;vertical-align:top;border:1px solid #e2e8f0;padding:4px 6px;text-align:left}"
           "th{background:#f1f5f9}td.en{color:#475569}table.sig td{border:none}.page{page-break-after:always}"
           ".page:last-child{page-break-after:avoid}.note{font-style:italic;font-size:7.5pt;color:#64748b}")
    return (f"<!DOCTYPE html><html lang='de'><head><meta charset='utf-8'><title>Mietvertrag – Homelyo</title>"
            f"<style>{css}</style></head><body>{''.join(body)}</body></html>")
